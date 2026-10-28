#!/usr/bin/env python3
"""Build literature-curated H1/H0 edge likelihood references for V1.

The source workbooks are retained verbatim under ``gui/calibration``.  This
builder converts only unambiguous, gene-resolved interactions into a compact
TSV reference catalog, constructs deterministic endpoint-matched background
pairs where a workbook has no experimental negatives, and records each V1
edge stream's source Bayes factor on those references.

The generated training table is consumed by :mod:`workflow_engine`; the GUI
never needs to import Excel files at analysis time.
"""

from __future__ import annotations

import argparse
import json
import random
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from workflow_engine import (
    EDGE_LIKELIHOOD_PROFILES,
    PROJECT_ROOT,
    default_configuration,
    edge_stream_factor_table,
    empirical_likelihood_calibration,
    load_registry,
)


CALIBRATION_ROOT = Path("gui/calibration/positive_cases_20261028")
SOURCE_DIRECTORY = CALIBRATION_ROOT / "source/Positive Cases"
REFERENCE_FILE = CALIBRATION_ROOT / "curated_edge_likelihood_references.tsv"
TRAINING_FILE = CALIBRATION_ROOT / "curated_edge_likelihood_training.tsv"
SUMMARY_FILE = CALIBRATION_ROOT / "curated_edge_likelihood_summary.json"
UNRESOLVED_FILE = CALIBRATION_ROOT / "unresolved_reference_rows.tsv"

PROFILE_STREAMS = {
    "physical_proximity": (
        "mpkccd_localization",
        "compartments_localization",
        "hpa_primary",
        "hpa_high_confidence",
        "imcd_basal_compartment_presence",
        "imcd_ddavp_compartment_presence",
    ),
    "physical_binding": (
        "alphapulldown_structural",
        "biogrid_physical_interaction",
    ),
    "directed_phosphorylation": ("kinase_predictor",),
    "regulatory_signaling": ("omnipath_core", "string_v12"),
    "messenger_binding": ("stitch_secondary_messenger",),
}

TRANSCRIPTION_PATTERN = re.compile(
    r"\b(transcri(?:pt|ption|ptional)|promoter|mrna|gene expression)\b",
    flags=re.IGNORECASE,
)
GENE_TOKEN_PATTERN = re.compile(r"[A-Za-z][A-Za-z0-9]+")
GENERIC_TOKENS = {
    "actin",
    "alpha",
    "ampk",
    "beta",
    "camp",
    "catalytic",
    "cd",
    "cell",
    "cells",
    "ccd",
    "ck2",
    "complex",
    "duct",
    "enac",
    "erk",
    "family",
    "gamma",
    "gene",
    "gsk3",
    "human",
    "imcd",
    "isoform",
    "kinase",
    "mouse",
    "omcd",
    "pka",
    "pkg",
    "pp1",
    "pp2a",
    "pp2b",
    "protein",
    "rat",
    "regulatory",
    "subunit",
    "unresolved",
}

ALIASES = {
    "trpp2": "Pkd2",
    "pc2": "Pkd2",
    "trpp1": "Pkd1",
    "pc1": "Pkd1",
    "mr": "Nr3c2",
    "chip": "Stub1",
    "hsc70": "Hspa8",
    "bip": "Hspa5",
    "grp78": "Hspa5",
    "sap97": "Dlg1",
    "akap150": "Akap5",
    "akap220": "Akap11",
    "akap18": "Akap7",
    "af9": "Mllt3",
    "af17": "Mllt6",
    "dot1a": "Dot1l",
    "beta1pix": "Arhgef7",
    "polaris": "Ift88",
}

MESSENGER_ALIASES = {
    "camp": "SM_CAMP",
    "3,5-cyclic adenosine monophosphate": "SM_CAMP",
    "3',5'-cyclic adenosine monophosphate": "SM_CAMP",
    "cgmp": "SM_CGMP",
    "3,5-cyclic guanosine monophosphate": "SM_CGMP",
    "3',5'-cyclic guanosine monophosphate": "SM_CGMP",
    "ca2+": "SM_CA2",
    "calcium": "SM_CA2",
    "ip3": "SM_IP3",
    "dag": "SM_DAG",
    "pip3": "SM_PIP3",
    "pip2": "SM_PIP2",
    "no": "SM_NO",
    "cadpr": "SM_CADPR",
    "naadp": "SM_NAADP",
    "adpr": "SM_ADPR",
    "h2o2": "SM_H2O2",
    "s1p": "SM_S1P",
}


def _clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _canonical_pair(left: str, right: str) -> tuple[str, str]:
    return tuple(sorted((str(left), str(right)), key=lambda item: item.casefold()))


def _worksheet_records(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        header_row = None
        headers: list[str] = []
        for row_index in range(1, min(worksheet.max_row, 20) + 1):
            row = [
                _clean_text(worksheet.cell(row_index, column).value)
                for column in range(1, worksheet.max_column + 1)
            ]
            if len([value for value in row if value]) >= 2 and any(
                value.casefold()
                in {
                    "source",
                    "regulator",
                    "molecule",
                    "evidence id",
                    "protein a / bait",
                }
                for value in row
            ):
                header_row = row_index
                headers = row
                break
        if header_row is None:
            raise ValueError(f"could not locate table header in {path.name}: {sheet_name}")
        records: list[dict[str, Any]] = []
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            values = [
                worksheet.cell(row_index, column).value
                for column in range(1, worksheet.max_column + 1)
            ]
            if not any(_clean_text(value) for value in values):
                continue
            record = {
                header: value
                for header, value in zip(headers, values)
                if header
            }
            record["_source_row"] = row_index
            records.append(record)
        return records
    finally:
        workbook.close()


def _add_symbols_from_table(
    catalog: dict[str, str],
    path: Path,
    *,
    columns: Iterable[str],
    compression: str | None = "infer",
) -> None:
    if not path.exists():
        return
    try:
        table = pd.read_csv(path, sep="\t", dtype=str, compression=compression).fillna("")
    except Exception:
        return
    for column in columns:
        if column not in table.columns:
            continue
        for symbol in table[column].astype(str):
            symbol = symbol.strip()
            if symbol and re.fullmatch(r"[A-Za-z][A-Za-z0-9._-]*", symbol):
                catalog.setdefault(symbol.casefold(), symbol)


def build_symbol_catalog(project: Path) -> dict[str, str]:
    catalog: dict[str, str] = {}
    paths = (
        (project / "data/node_selection/node_universe_combined_nonzero.tsv", ("symbol",)),
        (project / "data/node_selection/mouse_signaling_nodes_liberal.tsv", ("symbol",)),
        (
            project
            / "data/node_selection/collecting_duct/raw/ensembl116_rat_mouse_orthologs.tsv",
            ("rat_gene_symbol", "mouse_gene_symbol", "rat_symbol", "mouse_symbol"),
        ),
        (
            project
            / "data/edge_characterization/imcd_compartment_presence/processed/"
            "imcd_compartment_node_profiles.tsv",
            ("symbol",),
        ),
        (
            project
            / "data/edge_characterization/compartments/integrated/processed/"
            "node_compartments_profiles.tsv.gz",
            ("symbol",),
        ),
        (
            project
            / "data/edge_characterization/biogrid/5.0.261/"
            "mapped_graph_physical_pairs.tsv.gz",
            ("node_a", "node_b"),
        ),
        (
            project
            / "data/edge_characterization/omnipath/2026-07-30/raw/"
            "omnipath_mouse_core_post_translational.tsv",
            ("source_genesymbol", "target_genesymbol", "source", "target"),
        ),
    )
    for path, columns in paths:
        _add_symbols_from_table(catalog, path, columns=columns)
    for symbol in MESSENGER_ALIASES.values():
        catalog.setdefault(symbol.casefold(), symbol)
    for alias, symbol in ALIASES.items():
        catalog.setdefault(symbol.casefold(), symbol)
        ALIASES[alias] = symbol
    return catalog


def _symbols_from_label(label: Any, catalog: dict[str, str]) -> list[str]:
    text = _clean_text(label)
    symbols: list[str] = []
    for token in GENE_TOKEN_PATTERN.findall(text):
        key = token.casefold()
        if key in GENERIC_TOKENS:
            continue
        symbol = ALIASES.get(key) or catalog.get(key)
        if symbol and symbol not in symbols:
            symbols.append(symbol)
    return symbols


def _single_symbol(label: Any, catalog: dict[str, str]) -> str | None:
    symbols = _symbols_from_label(label, catalog)
    return symbols[0] if len(symbols) == 1 else None


def _record_text(record: dict[str, Any]) -> str:
    return " ".join(_clean_text(value) for key, value in record.items() if not key.startswith("_"))


def _reference_row(
    *,
    profile: str,
    node_a: str,
    node_b: str,
    label: str,
    context_policy: str,
    source_workbook: str,
    source_sheet: str,
    source_row: int,
    source_record_id: str,
    evidence_status: str,
    pmids: str,
    reasoning: str,
    negative_kind: str = "",
) -> dict[str, Any]:
    left, right = _canonical_pair(node_a, node_b)
    return {
        "profile": profile,
        "node_a": left,
        "node_b": right,
        "class_label": label,
        "context_policy": context_policy,
        "negative_kind": negative_kind,
        "source_workbook": source_workbook,
        "source_sheet": source_sheet,
        "source_row": int(source_row),
        "source_record_id": source_record_id,
        "evidence_status": evidence_status,
        "pmids": pmids,
        "reasoning": reasoning,
    }


def _physical_references(
    source_dir: Path,
    catalog: dict[str, str],
    unresolved: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    workbook = "Collecting_duct_PPI_evidence.xlsx"
    sheet = "Focused associations"
    rows: list[dict[str, Any]] = []
    for record in _worksheet_records(source_dir / workbook, sheet):
        if TRANSCRIPTION_PATTERN.search(_record_text(record)):
            continue
        left_symbols = _symbols_from_label(record.get("Protein A / bait"), catalog)
        right_symbols = _symbols_from_label(
            record.get("Protein B / partner or set"), catalog
        )
        if not left_symbols or not right_symbols:
            unresolved.append(
                {
                    "profile": "physical_proximity;physical_binding",
                    "source_workbook": workbook,
                    "source_sheet": sheet,
                    "source_row": record["_source_row"],
                    "source_record_id": _clean_text(record.get("Evidence ID")),
                    "node_a_label": _clean_text(record.get("Protein A / bait")),
                    "node_b_label": _clean_text(record.get("Protein B / partner or set")),
                    "reason": "no unambiguous gene-resolved endpoint",
                }
            )
            continue
        for node_a in left_symbols:
            for node_b in right_symbols:
                if node_a == node_b:
                    continue
                for profile in ("physical_proximity", "physical_binding"):
                    rows.append(
                        _reference_row(
                            profile=profile,
                            node_a=node_a,
                            node_b=node_b,
                            label="H1_POSITIVE",
                            context_policy="collecting_duct_specific",
                            source_workbook=workbook,
                            source_sheet=sheet,
                            source_row=record["_source_row"],
                            source_record_id=_clean_text(record.get("Evidence ID")),
                            evidence_status=_clean_text(record.get("Evidence level")),
                            pmids="",
                            reasoning=_clean_text(record.get("Limitations")),
                        )
                    )
    return rows


def _phosphorylation_references(
    source_dir: Path,
    catalog: dict[str, str],
    unresolved: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    workbook = "Collecting_duct_phosphorylation_evidence.xlsx"
    sheet = "Directed evidence"
    pair_records: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in _worksheet_records(source_dir / workbook, sheet):
        if _clean_text(record.get("Source type")).casefold() != "kinase":
            continue
        source = _single_symbol(record.get("source"), catalog)
        target = _single_symbol(record.get("target"), catalog)
        status = _clean_text(record.get("Evidence status"))
        if source is None or target is None or source == target:
            if status in {"H1_POSITIVE", "H0_EXPLICIT_NEGATIVE"}:
                unresolved.append(
                    {
                        "profile": "directed_phosphorylation",
                        "source_workbook": workbook,
                        "source_sheet": sheet,
                        "source_row": record["_source_row"],
                        "source_record_id": _clean_text(record.get("Evidence ID")),
                        "node_a_label": _clean_text(record.get("source")),
                        "node_b_label": _clean_text(record.get("target")),
                        "reason": "kinase/substrate endpoint is unresolved or non-gene",
                    }
                )
            continue
        if status == "H0_EXPLICIT_NEGATIVE" and "excluded from calibration h0" in _record_text(record).casefold():
            continue
        if status not in {"H1_POSITIVE", "H0_EXPLICIT_NEGATIVE"}:
            continue
        pair_records[(source, target)].append(record)

    rows: list[dict[str, Any]] = []
    for (source, target), records in pair_records.items():
        statuses = {_clean_text(record.get("Evidence status")) for record in records}
        label = (
            "H1_POSITIVE"
            if "H1_POSITIVE" in statuses
            else "H0_EXPLICIT_NEGATIVE"
        )
        chosen = next(
            record
            for record in records
            if _clean_text(record.get("Evidence status")) == label
        )
        rows.append(
            _reference_row(
                profile="directed_phosphorylation",
                node_a=source,
                node_b=target,
                label=label,
                context_policy="context_general_biochemical",
                source_workbook=workbook,
                source_sheet=sheet,
                source_row=chosen["_source_row"],
                source_record_id=";".join(
                    sorted(
                        {
                            _clean_text(record.get("Evidence ID"))
                            for record in records
                            if _clean_text(record.get("Evidence ID"))
                        }
                    )
                ),
                evidence_status=label,
                pmids=";".join(
                    sorted(
                        {
                            _clean_text(record.get("PMID"))
                            for record in records
                            if _clean_text(record.get("PMID"))
                        }
                    )
                ),
                reasoning=_clean_text(chosen.get("Qualification / limitations")),
                negative_kind=(
                    "explicit_tested_null" if label.startswith("H0") else ""
                ),
            )
        )
    return rows


def _regulatory_references(
    source_dir: Path,
    catalog: dict[str, str],
    unresolved: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    workbook = "collecting_duct_direct_regulation.xlsx"
    sheet = "Direct + CD support"
    rows: list[dict[str, Any]] = []
    for record in _worksheet_records(source_dir / workbook, sheet):
        if TRANSCRIPTION_PATTERN.search(_record_text(record)):
            continue
        source = _single_symbol(record.get("Regulator"), catalog)
        target = _single_symbol(record.get("Target"), catalog)
        if source is None or target is None or source == target:
            unresolved.append(
                {
                    "profile": "regulatory_signaling",
                    "source_workbook": workbook,
                    "source_sheet": sheet,
                    "source_row": record["_source_row"],
                    "source_record_id": _clean_text(record.get("ID")),
                    "node_a_label": _clean_text(record.get("Regulator")),
                    "node_b_label": _clean_text(record.get("Target")),
                    "reason": "regulator/target endpoint is unresolved or multi-protein",
                }
            )
            continue
        rows.append(
            _reference_row(
                profile="regulatory_signaling",
                node_a=source,
                node_b=target,
                label="H1_POSITIVE",
                context_policy="collecting_duct_specific",
                source_workbook=workbook,
                source_sheet=sheet,
                source_row=record["_source_row"],
                source_record_id=_clean_text(record.get("ID")),
                evidence_status="direct mechanism with collecting-duct support",
                pmids=_clean_text(record.get("PMID")),
                reasoning=_clean_text(record.get("Qualification")),
            )
        )
    return rows


def _messenger_symbol(value: Any) -> str | None:
    normalized = re.sub(r"\s+", " ", _clean_text(value)).strip().casefold()
    normalized = normalized.replace("′", "'").replace("’", "'")
    return MESSENGER_ALIASES.get(normalized)


def _messenger_references(
    source_dir: Path,
    catalog: dict[str, str],
    unresolved: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    workbook = "messenger_binding_review.xlsx"
    sheet = "Pair Classifications"
    rows: list[dict[str, Any]] = []
    for record in _worksheet_records(source_dir / workbook, sheet):
        molecule = _messenger_symbol(record.get("molecule"))
        protein = _single_symbol(record.get("protein"), catalog)
        label = _clean_text(record.get("final_label"))
        if label not in {"H1_POSITIVE", "H0_EXPLICIT_NEGATIVE"}:
            continue
        if molecule is None or protein is None:
            unresolved.append(
                {
                    "profile": "messenger_binding",
                    "source_workbook": workbook,
                    "source_sheet": sheet,
                    "source_row": record["_source_row"],
                    "source_record_id": "",
                    "node_a_label": _clean_text(record.get("molecule")),
                    "node_b_label": _clean_text(record.get("protein")),
                    "reason": "noncanonical messenger analog or unresolved protein",
                }
            )
            continue
        rows.append(
            _reference_row(
                profile="messenger_binding",
                node_a=molecule,
                node_b=protein,
                label=label,
                context_policy="context_general_biochemical",
                source_workbook=workbook,
                source_sheet=sheet,
                source_row=record["_source_row"],
                source_record_id="",
                evidence_status=label,
                pmids=_clean_text(record.get("PMIDs")),
                reasoning=_clean_text(record.get("reasoning")),
                negative_kind=(
                    "explicit_tested_null" if label.startswith("H0") else ""
                ),
            )
        )
    return rows


def _deduplicate_references(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_pair: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_pair[(row["profile"], row["node_a"], row["node_b"])].append(row)
    result: list[dict[str, Any]] = []
    for _, candidates in by_pair.items():
        positives = [row for row in candidates if row["class_label"].startswith("H1")]
        selected = positives or candidates
        first = dict(selected[0])
        first["source_record_id"] = ";".join(
            sorted(
                {
                    row["source_record_id"]
                    for row in selected
                    if row["source_record_id"]
                }
            )
        )
        first["pmids"] = ";".join(
            sorted({row["pmids"] for row in selected if row["pmids"]})
        )
        result.append(first)
    return sorted(
        result,
        key=lambda row: (
            row["profile"],
            row["class_label"],
            row["node_a"].casefold(),
            row["node_b"].casefold(),
        ),
    )


def _matched_background_rows(
    positives: list[dict[str, Any]],
    *,
    forbidden_pairs: set[tuple[str, str]],
    multiplier: int,
    seed: int,
) -> list[dict[str, Any]]:
    if not positives:
        return []
    target_count = len(positives) * multiplier
    left = [row["node_a"] for row in positives]
    right = [row["node_b"] for row in positives]
    endpoint_pool = sorted(set(left + right), key=str.casefold)
    rng = random.Random(seed)
    generated: list[dict[str, Any]] = []
    seen = set(forbidden_pairs)
    attempts = 0
    while len(generated) < target_count and attempts < target_count * 200:
        attempts += 1
        if attempts <= target_count * 20:
            node_a = left[rng.randrange(len(left))]
            node_b = right[rng.randrange(len(right))]
        else:
            node_a, node_b = rng.sample(endpoint_pool, 2)
        if node_a == node_b:
            continue
        pair = _canonical_pair(node_a, node_b)
        if pair in seen:
            continue
        seen.add(pair)
        template = positives[len(generated) % len(positives)]
        generated.append(
            _reference_row(
                profile=template["profile"],
                node_a=pair[0],
                node_b=pair[1],
                label="H0_BACKGROUND",
                context_policy=template["context_policy"],
                source_workbook="deterministic endpoint-matched background",
                source_sheet="generated",
                source_row=0,
                source_record_id=f"BG{len(generated) + 1:04d}",
                evidence_status="unlabeled background pair",
                pmids="",
                reasoning=(
                    "Endpoint-matched background; not an experimentally proven "
                    "noninteraction. Excludes every curated positive/explicit-negative "
                    "pair in this hypothesis class."
                ),
                negative_kind="matched_unlabeled_background",
            )
        )
    return generated


def build_reference_catalog(project: Path, source_dir: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    catalog = build_symbol_catalog(project)
    unresolved: list[dict[str, Any]] = []
    rows = (
        _physical_references(source_dir, catalog, unresolved)
        + _phosphorylation_references(source_dir, catalog, unresolved)
        + _regulatory_references(source_dir, catalog, unresolved)
        + _messenger_references(source_dir, catalog, unresolved)
    )
    rows = _deduplicate_references(rows)
    all_pairs_by_profile: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for row in rows:
        all_pairs_by_profile[row["profile"]].add((row["node_a"], row["node_b"]))
    for profile in ("physical_proximity", "physical_binding", "regulatory_signaling"):
        positives = [
            row
            for row in rows
            if row["profile"] == profile and row["class_label"].startswith("H1")
        ]
        rows.extend(
            _matched_background_rows(
                positives,
                forbidden_pairs=all_pairs_by_profile[profile],
                multiplier=5,
                seed={
                    "physical_proximity": 104729,
                    "physical_binding": 104759,
                    "regulatory_signaling": 104761,
                }[profile],
            )
        )
    messenger_rows = [row for row in rows if row["profile"] == "messenger_binding"]
    messenger_positives = [
        row for row in messenger_rows if row["class_label"].startswith("H1")
    ]
    messenger_negatives = [
        row for row in messenger_rows if row["class_label"].startswith("H0")
    ]
    if len(messenger_negatives) < 20:
        rows.extend(
            _matched_background_rows(
                messenger_positives,
                forbidden_pairs={
                    (row["node_a"], row["node_b"]) for row in messenger_rows
                },
                multiplier=5,
                seed=104773,
            )
        )
    references = pd.DataFrame(_deduplicate_references(rows))
    unresolved_frame = pd.DataFrame(unresolved)
    return references, unresolved_frame


def _factor_lookup(table: pd.DataFrame) -> dict[tuple[str, str], float]:
    lookup: dict[tuple[str, str], float] = defaultdict(lambda: 1.0)
    for row in table.itertuples(index=False):
        pair = _canonical_pair(str(row.node_a), str(row.node_b))
        lookup[pair] *= float(row.bayes_factor)
    return dict(lookup)


def build_training_table(
    project: Path,
    references: pd.DataFrame,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    registry = load_registry(project)
    configuration = default_configuration(registry)
    definitions = {stream["id"]: stream for stream in registry["edge_streams"]}
    training_rows: list[dict[str, Any]] = []
    models: dict[str, Any] = {}
    for profile, stream_ids in PROFILE_STREAMS.items():
        profile_references = references.loc[references["profile"].eq(profile)].copy()
        if profile_references.empty:
            continue
        symbols = sorted(
            set(profile_references["node_a"]).union(profile_references["node_b"]),
            key=str.casefold,
        )
        for stream_id in stream_ids:
            definition = definitions[stream_id]
            state = configuration["edge_streams"][stream_id]
            table = edge_stream_factor_table(project, definition, state, symbols)
            lookup = _factor_lookup(table)
            values = np.asarray(
                [
                    lookup.get(_canonical_pair(row.node_a, row.node_b), 1.0)
                    for row in profile_references.itertuples(index=False)
                ],
                dtype=float,
            )
            h1 = profile_references["class_label"].str.startswith("H1").to_numpy()
            h0 = profile_references["class_label"].str.startswith("H0").to_numpy()
            _, model = empirical_likelihood_calibration(
                values,
                h1_reference=h1,
                h0_reference=h0,
                evidence_observed=np.ones(len(values), dtype=bool),
                enabled=True,
                hypothesis_class=profile,
                anchor_description=(
                    "Primary-literature curation dated 2026-10-28. "
                    + (
                        "Collecting-duct specificity required. "
                        if profile_references["context_policy"].eq(
                            "collecting_duct_specific"
                        ).all()
                        else "Biochemical mechanism accepted across model contexts. "
                    )
                    + "H0 uses explicit tested nulls when available; otherwise "
                    "deterministic endpoint-matched unlabeled background pairs."
                ),
                smoothing=0.5,
                minimum_reference_count=20,
            )
            model["reference_source"] = "positive_cases_20261028"
            model["reference_file"] = str(REFERENCE_FILE).replace("\\", "/")
            model["context_policy"] = sorted(
                set(profile_references["context_policy"].astype(str))
            )
            model["explicit_h0_count"] = int(
                profile_references["negative_kind"].eq("explicit_tested_null").sum()
            )
            model["background_h0_count"] = int(
                profile_references["negative_kind"]
                .eq("matched_unlabeled_background")
                .sum()
            )
            models[stream_id] = model
            for reference, source_factor in zip(
                profile_references.to_dict(orient="records"), values
            ):
                training_rows.append(
                    {
                        "stream_id": stream_id,
                        **reference,
                        "source_bayes_factor": float(source_factor),
                        "source_observed": bool(
                            not np.isclose(source_factor, 1.0, atol=1e-12, rtol=0.0)
                        ),
                    }
                )
    training = pd.DataFrame(training_rows)
    summary = {
        "schema_version": 1,
        "curation_date": "2026-10-28",
        "source_package": "Positive Cases.zip supplied 2026-10-28",
        "reference_file": str(REFERENCE_FILE).replace("\\", "/"),
        "training_file": str(TRAINING_FILE).replace("\\", "/"),
        "reference_pair_count": int(len(references)),
        "profile_counts": {
            profile: {
                "h1": int(group["class_label"].str.startswith("H1").sum()),
                "h0": int(group["class_label"].str.startswith("H0").sum()),
                "explicit_h0": int(
                    group["negative_kind"].eq("explicit_tested_null").sum()
                ),
                "background_h0": int(
                    group["negative_kind"].eq("matched_unlabeled_background").sum()
                ),
                "context_policy": sorted(set(group["context_policy"].astype(str))),
            }
            for profile, group in references.groupby("profile", sort=True)
        },
        "stream_models_default_settings": models,
        "policy_notes": [
            "KinasePredictor and messenger binding use context-general biochemical references.",
            "Localization, structural/physical, OmniPath, and STRING references require collecting-duct support.",
            "STRING is calibrated as acute regulatory signaling, not generic functional association.",
            "Transcription/promoter-mediated records are excluded.",
            "A background H0 is an unlabeled comparison set, not proof of noninteraction.",
            "Derived closure streams retain their documented fixed source factor and are not independently recalibrated.",
        ],
    }
    return training, summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--source-dir", type=Path)
    args = parser.parse_args()
    project = args.project.resolve()
    source_dir = (
        args.source_dir.resolve()
        if args.source_dir
        else (project / SOURCE_DIRECTORY).resolve()
    )
    missing = [
        name
        for name in (
            "Collecting_duct_PPI_evidence.xlsx",
            "Collecting_duct_phosphorylation_evidence.xlsx",
            "collecting_duct_direct_regulation.xlsx",
            "messenger_binding_review.xlsx",
        )
        if not (source_dir / name).exists()
    ]
    if missing:
        raise FileNotFoundError("missing curated workbook(s): " + ", ".join(missing))
    references, unresolved = build_reference_catalog(project, source_dir)
    training, summary = build_training_table(project, references)
    output_root = project / CALIBRATION_ROOT
    output_root.mkdir(parents=True, exist_ok=True)
    references.to_csv(project / REFERENCE_FILE, sep="\t", index=False)
    training.to_csv(project / TRAINING_FILE, sep="\t", index=False)
    unresolved.to_csv(project / UNRESOLVED_FILE, sep="\t", index=False)
    (project / SUMMARY_FILE).write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
